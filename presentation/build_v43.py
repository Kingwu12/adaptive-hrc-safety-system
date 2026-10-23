"""Build FYP-assessment-presentation-v4.3.pptx from v4.2.

v4.3 (23 Oct 2026), from team feedback:
- a quick Acknowledgement of Country opens Zenan's script (unit staff on Ed
  #112: "a quick one would suffice");
- slide 2 shows the task on a real site beside our lab version of it, both
  photos large enough to read from the back of the room. Photo:
  "Drywalling the Ceiling" by Qole Pejorian, CC BY 2.0,
  https://www.flickr.com/photos/36824782@N00/34589819 (credited on the slide);
- bigger images on slides 4 and 6: the console is recaptured at a narrow
  window on its ready card (presentation/media/dashboard-console-ready.png),
  the model plot is redrawn at its size on the slide, and the rig slide shows
  one large photo beside the video in place of two small ones, with the
  OptiTrack camera, the Xsens suit and the UR10 arm circled and labelled on it;
- carbon is off slide 11 and out of its script. The cost and carbon answer
  stays in the guide's prepared questions;
- two edits the team made in the Drive copy of v4.2 are carried over: the
  slide 1 subtitle is the paper's title, and slide 8 says the 3 · 5 · 3 picks
  were "perceived equally as safe".

Zenan's script stays at about 400 words: the opening and slide 2 grew, slides
3 and 4 were tightened.

    python3 presentation/build_v43.py
"""
import copy
import math

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from build_v4 import BLUE, GREY, HERE, NAVY, ROOT, WHITE, box, notes, text
from build_v42 import shape

AMBER = RGBColor(0xFF, 0xB8, 0x1C)

SRC = HERE / "FYP-assessment-presentation-v4.2.pptx"
OUT = HERE / "FYP-assessment-presentation-v4.3.pptx"
MEDIA = HERE / "media"


def remove(sh):
    sh._element.getparent().remove(sh._element)


def add_line(run, line):
    """Add a line break after run, then line in run's formatting."""
    r = copy.deepcopy(run._r)
    r.find(qn("a:t")).text = line
    br = run._r.makeelement(qn("a:br"), {})
    run._r.addnext(br)
    br.addnext(r)


def callout(s, label, x, y, r, lx, ly, lw):
    """Circle (x, y) with radius r and point at it from a dark label whose
    top left is (lx, ly); all in inches. A newline in label breaks the line."""
    lines = label.split("\n")
    h = 0.16 + 0.24 * len(lines)
    ring = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - r), Inches(y - r), Inches(2 * r), Inches(2 * r))
    ring.fill.background()
    ring.line.color.rgb = AMBER
    ring.line.width = Pt(3)
    ring.shadow.inherit = False
    # the arrow runs from the label's edge to the circle's edge
    cx, cy = lx + lw / 2, ly + h / 2
    dx, dy = x - cx, y - cy
    d = math.hypot(dx, dy)
    t = min(lw / 2 / abs(dx) if dx else math.inf, h / 2 / abs(dy) if dy else math.inf)
    line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(cx + dx * t), Inches(cy + dy * t),
                                  Inches(x - dx / d * r), Inches(y - dy / d * r))
    line.line.color.rgb = AMBER
    line.line.width = Pt(2.5)
    line.line._ln.append(line.line._ln.makeelement(qn("a:tailEnd"), {"type": "triangle"}))
    box(s, lx, ly, lw, h, fill=NAVY)
    tf = text(s, lines[0], lx, ly, lw, h, size=15, bold=True, color=WHITE,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE).text_frame
    tf.word_wrap = False
    for line in lines[1:]:
        add_line(tf.paragraphs[0].runs[-1], line)


def crop(pic, left, top, right, bottom, x, y, w):
    """Crop by source pixels (left, top, right, bottom) and size the frame to
    the cropped aspect at width w."""
    pw, ph = pic.image.size
    pic.crop_left, pic.crop_right = left / pw, (pw - right) / pw
    pic.crop_top, pic.crop_bottom = top / ph, (ph - bottom) / ph
    pic.left, pic.top, pic.width = Inches(x), Inches(y), Inches(w)
    pic.height = Inches(w * (bottom - top) / (right - left))


def build():
    prs = Presentation(str(SRC))
    s1, s2, s3, s4, s5, s6, s7, s8, s9, s10, s11, s12 = prs.slides

    # Two wording edits the team made in the Drive copy of v4.2 (21 Oct):
    # the subtitle is the paper's title, broken where it balances, and the
    # 3 · 5 · 3 caption
    sub = shape(s1, "Human–robot ceiling-panel installation").text_frame
    sub.word_wrap = True
    run = sub.paragraphs[0].runs[0]
    run.text = "Safety and Efficiency of Context-Aware Speed and Separation Monitoring"
    run.font.size = Pt(20)
    add_line(run, "in Human–Robot Ceiling-Panel Installation")
    shape(s8, "TextBox 7").text_frame.paragraphs[0].runs[0].text = (
        "times each controller was picked as ‘safest’: perceived equally as safe")
    shape(s8, "TextBox 9").text_frame.paragraphs[0].runs[0].text = (
        "median time for the robot to stop (max 0.23 s)")

    # 02: the task on a real site beside our lab version, both large and
    # cropped square to the people and the panel; the three points narrow to
    # the right, spaced by their line counts (3, 4, 3), and the question spans
    # the bottom above the footer band
    remove(shape(s2, "Picture 3"))
    for path, (l, t, r, b), x, head, body in (
            (MEDIA / "site-ceiling-board.jpg", (284, 60, 990, 766), 0.6,
             "On site today", "A second worker holds the board."),
            (ROOT / "paper/figures/task-frame.jpg", (80, 0, 800, 720), 4.3,
             "In our lab", "The robot holds the panel.")):
        crop(s2.shapes.add_picture(str(path), 0, 0), l, t, r, b, x, 1.3, 3.5)
        text(s2, [[(head, {"bold": True, "color": BLUE})], body], x, 4.88, 3.5, 0.7, size=15)
    text(s2, "Photo: Qole Pejorian, CC BY 2.0", 0.6, 4.96, 3.5, 0.22, size=9, color=GREY, align=PP_ALIGN.RIGHT)
    bodies = ("A robot can take the weight.",
              "Safety zones stop it whenever someone is close.",
              "Crews need it right both ways.")
    for i, (body, y) in enumerate(zip(bodies, (1.4, 2.82, 4.51))):
        oval, tb = shape(s2, f"Oval {4 + 2 * i}"), shape(s2, f"TextBox {5 + 2 * i}")
        oval.left, oval.top = Inches(8.4), Inches(y + 0.03)
        tb.left, tb.top, tb.width = Inches(9.05), Inches(y), Inches(3.65)
        tb.text_frame.paragraphs[1].runs[0].text = body
        for p in tb.text_frame.paragraphs:
            p.runs[0].font.size = Pt(16)
    stop = shape(s2, "TextBox 9").text_frame.paragraphs[0].runs[0]
    stop.text = "Stop too late: injury."
    add_line(stop, "Stop too often: no one uses it.")
    q, qt = shape(s2, "Rounded Rectangle 10"), shape(s2, "TextBox 11")
    q.left, q.top, q.width, q.height = Inches(0.6), Inches(5.66), Inches(12.2), Inches(0.76)
    qt.left, qt.top, qt.width, qt.height = Inches(0.75), Inches(5.68), Inches(11.9), Inches(0.72)
    for r in qt.text_frame.paragraphs[0].runs:
        r.font.size = Pt(18)

    # 04: both pictures 3.85 in tall. The console is recaptured at a narrow
    # window so its text is large, showing its ready card as in a session
    # (the old capture was taken with no lab hardware and read "0/5 systems
    # ready"); a stand-in data feed supplied the ready state. The model plot is
    # redrawn at its size on the slide (scripts/make_deck_model_figures.py)
    remove(shape(s4, "Picture 2"))
    console = s4.shapes.add_picture(str(MEDIA / "dashboard-console-ready.png"),
                                    Inches(0.6), Inches(1.3), height=Inches(3.85))
    console.line.color.rgb = GREY
    console.line.width = Pt(0.75)
    remove(shape(s4, "Picture 3"))
    s4.shapes.add_picture(str(MEDIA / "model/held-out-trial-slide.png"),
                          Inches(7.15), Inches(1.3), Inches(5.65), Inches(3.85))
    for name, x, w in (("TextBox 4", 0.6, console.width / 914400), ("TextBox 5", 7.15, 5.65)):
        tb = shape(s4, name)
        tb.left, tb.top, tb.width = Inches(x), Inches(5.22), Inches(w)
    shape(s4, "TextBox 5").text_frame.paragraphs[1].runs[0].text = (
        "Upper bar: what they were doing. Lower bar: what the model said, live.")
    shape(s4, "TextBox 6").top = Inches(6.2)

    # 06: one large photo of the rig beside the video, in place of two small
    # ones (the top one was a back view that repeated the video). The play
    # label moves to the video's bottom right, off the participant's head, and
    # the captions are reworded so no word sits alone on a line
    remove(shape(s6, "Picture 4"))
    crop(shape(s6, "Picture 3"), 340, 0, 927, 675, 8.45, 1.35, 4.4)
    for name, x, y in (("Rounded Rectangle 11", 5.18, 4.92), ("TextBox 12", 5.24, 4.95)):
        shape(s6, name).left, shape(s6, name).top = Inches(x), Inches(y)
    for name, body in (("TextBox 7", "head position to about one millimetre"),
                       ("TextBox 8", "tracks the motion of all 23 body segments")):
        shape(s6, name).text_frame.paragraphs[1].runs[0].text = body
    # The hardware the captions name, circled on the photo (King, 23 Oct).
    # Positions are photo pixels in apparatus-a.jpg: the camera on the
    # overhead frame (clear of the silver box hanging beside it), the suit on
    # the torso, and the arm's upper end below the panel, clear of the suit's
    # circle and the head cap. Labels sit on plain ceiling, glass and wall.
    # 587 px of photo span 4.4 in
    k = 4.4 / 587
    for label, (px, py, pr), (lx, ly, lw) in (
            ("OptiTrack\ncamera", (655, 75, 20), (8.55, 1.45, 1.15)),
            ("Xsens suit", (690, 385, 52), (8.55, 3.95, 1.2)),
            ("UR10\nrobot arm", (730, 219, 40), (11.78, 3.55, 1.04))):
        callout(s6, label, 8.45 + (px - 340) * k, 1.35 + py * k, pr * k, lx, ly, lw)

    # 11: five cards, carbon removed; the bottom two take the full width
    for sh in list(s11.shapes):
        if Inches(1.3) < sh.top < Inches(6.8):
            remove(sh)
    cards = [
        ("Safety", "A research prototype, not a certified safety function. Allow 0.5 s to react, not 0.4."),
        ("Whole-life cost", "Cameras, body suits, calibration, maintenance and training for every crew."),
        ("Worker health", "Removes overhead lifting strain; the suit itself was rated low restriction (2 of 5)."),
        ("People and trust", "Participants rated all three as safe. Crews must accept a robot as a co-worker."),
        ("Standards", "ISO 10218 and ISO/TS 15066 set the bar any site system must meet."),
    ]
    for i, (h, b) in enumerate(cards):
        if i < 3:
            x, y, w = 0.6 + i * 4.15, 1.4, 3.9
        else:
            x, y, w = 0.6 + (i - 3) * 6.225, 3.95, 5.975
        box(s11, x, y, w, 2.35)
        text(s11, [[(h, {"bold": True, "color": BLUE, "size": 21})], b], x + 0.2, y + 0.15, w - 0.4, 2.15, size=18)

    for i, body in SCRIPTS.items():
        notes(prs.slides[i - 1], body)
    prs.save(str(OUT))
    print("wrote", OUT)


SCRIPTS = {
1: """ZENAN | 0:30

Good afternoon. We acknowledge the Traditional Owners of the land we are meeting on, and pay our respects to their Elders.

I am Zenan, and with Luke and Michael I will present our project: a context-aware safety system that lets a collaborative robot recognise what a construction worker is doing, so it can keep them safe without stopping unnecessarily. I will introduce the problem and our system, Michael will present our testing and results, and Luke will cover project management and use on site.""",

2: """ZENAN | 0:50

Installing ceiling panels requires workers to hold heavy boards overhead. On a typical site, on the left, one person simply holds the board while another fixes it. In our laboratory, beside it, a collaborative robot does the holding.

However, the worker now stands beside a moving robot. Current practice defines fixed safety zones around the robot and stops it whenever a person is inside. These zones cannot distinguish a worker approaching to begin a task from one moving away. If the robot stops too late, the worker may be injured; if it stops too often, crews are likely to disable it. Solving this benefits installation crews, contractors and the engineers who certify these systems.

This leads to our research question, shown in the dark box: can a robot that is aware of the task phase keep workers as safe as fixed zones, while maintaining task efficiency?""",

3: """ZENAN | 0:50

To answer this, we compared three controllers, each adding one source of information. The fixed-zone controller uses separation distance only; this is current practice and our baseline. The reactive controller also considers approach speed. Our predictive controller adds the task phase, whether the worker is approaching, working or moving away, together with a predictor that anticipates entry into the stop zone.

The key design principle is shown in the dark bar: the learned components can only reduce the robot's speed, never increase it, so safety never depends on the machine-learning model being correct.""",

4: """ZENAN | 0:50

My contribution was the software. On the left is the operator console we developed to run each session. It verifies every sensor, assigns the controller, and records both the robot and the participant.

On the right is our trained phase model, tested on a participant it never saw. The upper bar shows what the participant was actually doing, and the lower bar shows the model's live prediction. It tracks the long working phase closely and is least certain at the transitions.

Michael will now present how we tested the system and what we found.""",

11: """LUKE | 0:45

Moving from the laboratory to a site introduces real constraints. For safety, this is a research prototype running on the robot's standard speed controls, not a certified safety function, and we would allow half a second to react. For cost, the cameras and suits require calibration, maintenance and training for every crew, weighed against the labour saved. For health, it removes overhead lifting from workers, and participants rated the suit as only mildly restrictive. Finally, crews must accept a robotic co-worker, and any site system must meet the relevant robot safety standards.""",
}


if __name__ == "__main__":
    build()
